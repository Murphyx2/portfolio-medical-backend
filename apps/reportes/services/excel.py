"""Excel/ZIP builders for Servicios prestados + Paquete ARS
(Requirements/ReportPage/REPORTES_REQUIREMENTS.md sections 5/6)."""

import io
import zipfile

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from apps.reportes.services.engine import (
    ServiceLine,
    ars_program_label,
    distinct_ars_program_slices,
    month_label,
    servicios_prestados_rows,
)

HEADER_FILL = PatternFill(start_color="FFE7F6EC", end_color="FFE7F6EC", fill_type="solid")
HEADER_FONT = Font(bold=True)
MONEY_FORMAT = "#,##0.00"
COLUMNS = ["Cant", "Descripción", "Precio RD$", "Valor RD$"]


def _safe_filename_part(text: str) -> str:
    return "".join(c for c in text if c not in '\\/:*?"<>|').strip()


def build_servicios_prestados_workbook(
    *, rows: list[ServiceLine], label: str, year: int, month: int, centro_name: str | None, user
) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Servicios prestados"

    ws.append(["Servicios prestados"])
    ws.append([label])
    ws.append([month_label(year, month)])
    ws.append([centro_name or "Todos"])
    who = getattr(user, "get_full_name", lambda: "")() or getattr(user, "username", "")
    ws.append([f"Generado {timezone.now():%Y-%m-%d %H:%M} por {who}"])
    ws.append([])

    header_row = ws.max_row + 1
    ws.append(COLUMNS)
    for cell in ws[header_row]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    total_valor = 0
    for row in rows:
        ws.append([row.cant, row.name, float(row.precio), float(row.valor)])
        for cell in ws[ws.max_row]:
            if cell.column_letter in ("C", "D"):
                cell.number_format = MONEY_FORMAT
        total_valor += row.valor

    ws.append(["", "", "Total", float(total_valor)])
    total_cell = ws.cell(row=ws.max_row, column=4)
    total_cell.number_format = MONEY_FORMAT
    ws.cell(row=ws.max_row, column=3).font = HEADER_FONT

    for col, width in zip("ABCD", (10, 40, 16, 16)):
        ws.column_dimensions[col].width = width

    return wb


def build_individual_filename(*, ars, programa, year: int, month: int) -> str:
    label = ars_program_label(ars, programa).replace(" - ", " - ")
    return f"{_safe_filename_part(label)} - {year:04d}-{month:02d}.xlsx"


def build_pack_zip(*, year: int, month: int, centro_id: int | None, centro_name: str | None, user) -> bytes:
    slices = distinct_ars_program_slices(year=year, month=month, centro_id=centro_id)
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for ars, programa in slices:
            rows = servicios_prestados_rows(
                year=year,
                month=month,
                ars_id=ars.id if ars else None,
                programa_id=programa.id if programa else None,
                centro_id=centro_id,
            )
            if not rows:
                continue
            label = ars_program_label(ars, programa)
            wb = build_servicios_prestados_workbook(
                rows=rows, label=label, year=year, month=month, centro_name=centro_name, user=user
            )
            xlsx_buffer = io.BytesIO()
            wb.save(xlsx_buffer)
            filename = build_individual_filename(ars=ars, programa=programa, year=year, month=month)
            zf.writestr(filename, xlsx_buffer.getvalue())
    return buffer.getvalue()
